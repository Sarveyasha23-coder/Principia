"""
Principia — local-first knowledge workspace
Run:
    pip install -r requirements.txt
    streamlit run app.py

All uploaded content and the SQLite database stay on this machine.
The sentence-transformer model may download once on first use unless already cached.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "principia.sqlite3"
MODEL_NAME = "all-MiniLM-L6-v2"
CHUNK_WORDS = 220
OVERLAP_WORDS = 35

st.set_page_config(
    page_title="Principia — Knowledge, connected",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
:root { --ink:#eef2ff; --muted:#a7b0ca; --line:rgba(255,255,255,.10); }
.stApp { background: radial-gradient(ellipse at 10% 0%, #202a53 0%, #101629 43%, #0b1020 100%); color:var(--ink); font-family:Inter,sans-serif; }
[data-testid="stSidebar"] { background:rgba(7,11,24,.88); border-right:1px solid var(--line); }
.block-container { padding-top:1.5rem; max-width:1450px; }
.hero { padding:1.7rem 1.8rem; border:1px solid rgba(167,139,250,.30); border-radius:22px; background:linear-gradient(115deg,rgba(79,70,229,.28),rgba(139,92,246,.12)); margin-bottom:1.2rem; }
.eyebrow { color:#c4b5fd; text-transform:uppercase; letter-spacing:.16em; font-size:.72rem; font-weight:800; }
.hero h1 { font-size:2.4rem; letter-spacing:-.055em; margin:.25rem 0; color:#fff; }
.hero p { color:#c3cbe0; margin:0; font-size:1rem; }
.panel { border:1px solid var(--line); background:rgba(255,255,255,.035); padding:1rem 1.1rem; border-radius:16px; }
div[data-testid="stMetric"] { background:rgba(255,255,255,.04); border:1px solid var(--line); padding:14px 16px; border-radius:14px; }
div[data-testid="stMetricLabel"] { color:#a7b0ca; }
.stButton button, .stDownloadButton button { border-radius:10px; border:1px solid rgba(167,139,250,.3); font-weight:650; }
.stButton button[kind="primary"] { background:linear-gradient(90deg,#4f46e5,#8b5cf6); color:white; border:0; }
.stTextInput input, .stTextArea textarea { background:rgba(255,255,255,.045); color:#fff; border-color:var(--line); border-radius:10px; }
div[data-testid="stFileUploader"] { border:1px dashed rgba(167,139,250,.4); border-radius:14px; padding:.5rem; }
div[data-testid="stExpander"] { border:1px solid var(--line); border-radius:12px; }
.small-muted { color:#a7b0ca; font-size:.88rem; }
</style>
""", unsafe_allow_html=True)


# ----------------------------- Database ------------------------------------
def connect():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    with connect() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS brains (
            name TEXT PRIMARY KEY,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            brain TEXT NOT NULL,
            source TEXT NOT NULL,
            content_hash TEXT NOT NULL,
            added_at TEXT NOT NULL,
            tags TEXT NOT NULL DEFAULT '[]',
            char_count INTEGER NOT NULL DEFAULT 0,
            UNIQUE(brain, content_hash),
            FOREIGN KEY(brain) REFERENCES brains(name) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            brain TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            text TEXT NOT NULL,
            embedding BLOB,
            embedding_dim INTEGER,
            FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_chunks_brain ON chunks(brain);
        CREATE INDEX IF NOT EXISTS idx_docs_brain ON documents(brain);
        """)
        c.execute(
            "INSERT OR IGNORE INTO brains(name, created_at) VALUES (?, ?)",
            ("default", datetime.now().isoformat(timespec="seconds")),
        )


@st.cache_resource(show_spinner="Loading local embedding model…")
def get_model():
    return SentenceTransformer(MODEL_NAME)


def list_brains():
    with connect() as c:
        return [r["name"] for r in c.execute("SELECT name FROM brains ORDER BY name")]


def valid_brain_name(name: str) -> str:
    name = re.sub(r"[^a-zA-Z0-9 _-]", "", name).strip()
    name = re.sub(r"\s+", "-", name)[:48]
    if not name:
        raise ValueError("Use letters, numbers, spaces, hyphens, or underscores.")
    return name


def create_brain(name: str):
    name = valid_brain_name(name)
    with connect() as c:
        c.execute(
            "INSERT OR IGNORE INTO brains(name, created_at) VALUES (?, ?)",
            (name, datetime.now().isoformat(timespec="seconds")),
        )
    return name


# ----------------------------- Ingestion -----------------------------------
def extract_text(upload) -> str:
    name = upload.name.lower()
    raw = upload.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n\n".join(parts).strip()
    return raw.decode("utf-8-sig", errors="replace").strip()


def chunk_text(text: str, size: int = CHUNK_WORDS, overlap: int = OVERLAP_WORDS) -> list[str]:
    words = re.findall(r"\S+", text)
    if not words:
        return []
    chunks, start = [], 0
    while start < len(words):
        end = min(start + size, len(words))
        chunk = " ".join(words[start:end]).strip()
        if chunk:
            chunks.append(chunk)
        if end == len(words):
            break
        start = max(end - overlap, start + 1)
    return chunks


def ingest_file(brain: str, upload, tags: list[str], model) -> tuple[str, int]:
    text = extract_text(upload)
    if not text:
        return "empty", 0
    digest = hashlib.sha256(upload.getvalue()).hexdigest()
    chunks = chunk_text(text)
    if not chunks:
        return "empty", 0

    with connect() as c:
        existing = c.execute(
            "SELECT id FROM documents WHERE brain=? AND content_hash=?",
            (brain, digest),
        ).fetchone()
        if existing:
            return "duplicate", 0

        cur = c.execute("""
            INSERT INTO documents(brain, source, content_hash, added_at, tags, char_count)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (brain, upload.name, digest, datetime.now().isoformat(timespec="seconds"),
              json.dumps(tags), len(text)))
        doc_id = cur.lastrowid

    # Batch encoding is much faster than one model call per chunk.
    embeddings = model.encode(chunks, normalize_embeddings=True, show_progress_bar=False)
    with connect() as c:
        c.executemany("""
            INSERT INTO chunks(document_id, brain, chunk_index, text, embedding, embedding_dim)
            VALUES (?, ?, ?, ?, ?, ?)
        """, [
            (doc_id, brain, i, chunk, np.asarray(emb, dtype=np.float32).tobytes(), len(emb))
            for i, (chunk, emb) in enumerate(zip(chunks, embeddings))
        ])
    return "added", len(chunks)


# ----------------------------- Retrieval -----------------------------------
def get_chunks(brain: str):
    with connect() as c:
        rows = c.execute("""
            SELECT ch.id, ch.text, ch.embedding, ch.embedding_dim,
                   d.source, d.added_at, d.tags, d.id AS document_id
            FROM chunks ch JOIN documents d ON d.id=ch.document_id
            WHERE ch.brain=? ORDER BY ch.id
        """, (brain,)).fetchall()
    return [dict(r) for r in rows]


def retrieve(query: str, brain: str, model, mode: str, top_k: int):
    rows = get_chunks(brain)
    if not query.strip() or not rows:
        return []
    texts = [r["text"] for r in rows]
    scores = np.zeros(len(rows), dtype=float)

    if mode in ("Hybrid", "Semantic"):
        q = model.encode([query], normalize_embeddings=True)[0].astype(np.float32)
        for i, row in enumerate(rows):
            blob, dim = row["embedding"], row["embedding_dim"]
            if blob and dim == len(q):
                emb = np.frombuffer(blob, dtype=np.float32)
                if len(emb) == len(q) and np.isfinite(emb).all():
                    scores[i] = float(np.dot(q, emb))
        if mode == "Semantic":
            pass
        else:
            scores *= 0.70

    if mode in ("Hybrid", "Keyword"):
        try:
            vectorizer = TfidfVectorizer(ngram_range=(1, 2), stop_words="english")
            matrix = vectorizer.fit_transform(texts)
            qvec = vectorizer.transform([query])
            keyword_scores = cosine_similarity(qvec, matrix)[0]
        except ValueError:
            keyword_scores = np.zeros(len(rows), dtype=float)
        if mode == "Keyword":
            scores = keyword_scores
        else:
            scores += 0.30 * keyword_scores

    indices = np.argsort(scores)[::-1][:top_k]
    results = []
    for idx in indices:
        if scores[idx] <= 0:
            continue
        row = rows[int(idx)]
        row["score"] = float(scores[idx])
        results.append(row)
    return results


def answer_from_results(question: str, results: list[dict]) -> str:
    if not results:
        return "I couldn't find a relevant passage in this brain. Try different wording or add a document that covers this topic."
    # Extractive, transparent answer: no claim that a generative model was used.
    terms = set(re.findall(r"\b[a-zA-Z0-9]{3,}\b", question.lower()))
    candidates = []
    for result in results[:5]:
        for sentence in re.split(r"(?<=[.!?])\s+", result["text"]):
            sentence = sentence.strip()
            if len(sentence) < 25:
                continue
            st_terms = set(re.findall(r"\b[a-zA-Z0-9]{3,}\b", sentence.lower()))
            overlap = len(terms & st_terms) / max(1, len(terms))
            candidates.append((overlap + min(len(sentence), 300) / 3000, sentence, result))
    candidates.sort(key=lambda x: x[0], reverse=True)
    selected, seen = [], set()
    for _, sentence, result in candidates:
        key = sentence.lower()
        if key in seen:
            continue
        seen.add(key)
        selected.append((sentence, result))
        if len(selected) == 3:
            break
    if not selected:
        selected = [(r["text"][:500], r) for r in results[:2]]
    lines = ["Here are the most relevant passages I found in your documents:"]
    for sentence, result in selected:
        lines.append(f'- “{sentence}” — **{result["source"]}** (relevance {result["score"]:.2f})')
    lines.append("\nThis is an extractive answer based on retrieved text, not a general-purpose generative AI response.")
    return "\n".join(lines)


# ----------------------------- UI ------------------------------------------
def get_stats(brain: str):
    with connect() as c:
        docs = c.execute("SELECT COUNT(*) n FROM documents WHERE brain=?", (brain,)).fetchone()["n"]
        chunks = c.execute("SELECT COUNT(*) n FROM chunks WHERE brain=?", (brain,)).fetchone()["n"]
        tags_rows = c.execute("SELECT tags FROM documents WHERE brain=?", (brain,)).fetchall()
        oldest = c.execute("SELECT MIN(added_at) v FROM documents WHERE brain=?", (brain,)).fetchone()["v"]
        source_rows = c.execute("SELECT source, added_at, tags, char_count FROM documents WHERE brain=? ORDER BY added_at DESC", (brain,)).fetchall()
    tags = set()
    for row in tags_rows:
        try:
            tags.update(json.loads(row["tags"]))
        except (ValueError, TypeError):
            pass
    return {"documents": docs, "chunks": chunks, "tags": len(tags), "created": oldest, "sources": [dict(r) for r in source_rows]}


def main():
    init_db()
    if "chat_history" not in st.session_state:
        st.session_state.chat_history = {}
    if "selected_brain" not in st.session_state:
        st.session_state.selected_brain = "default"

    with st.sidebar:
        st.markdown("## 🧠 Principia")
        st.caption("Your local knowledge workspace")
        brains = list_brains()
        current = st.session_state.selected_brain
        if current not in brains:
            current = "default"
        selected = st.selectbox("Workspace", brains, index=brains.index(current))
        st.session_state.selected_brain = selected
        st.divider()
        with st.expander("＋ Create workspace"):
            new_name = st.text_input("Workspace name", key="new_brain_name", placeholder="e.g. product-research")
            if st.button("Create workspace", use_container_width=True):
                try:
                    created = create_brain(new_name)
                    st.session_state.selected_brain = created
                    st.success(f"Workspace ready: {created}")
                    st.rerun()
                except ValueError as e:
                    st.error(str(e))
        st.divider()
        st.markdown("### Add knowledge")
        uploads = st.file_uploader("PDF, Word, TXT, or Markdown", type=["pdf", "docx", "txt", "md"], accept_multiple_files=True)
        tags_raw = st.text_input("Tags", placeholder="research, product, strategy")
        tags = list(dict.fromkeys(t.strip() for t in tags_raw.split(",") if t.strip()))
        if st.button("Process documents", type="primary", use_container_width=True):
            if not uploads:
                st.warning("Choose at least one file first.")
            else:
                try:
                    model = get_model()
                    added, duplicates, empty = 0, 0, 0
                    progress = st.progress(0)
                    status = st.empty()
                    for i, upload in enumerate(uploads):
                        status.write(f"Processing {upload.name}…")
                        result, count = ingest_file(selected, upload, tags, model)
                        if result == "added":
                            added += 1
                        elif result == "duplicate":
                            duplicates += 1
                        else:
                            empty += 1
                        progress.progress((i + 1) / len(uploads))
                    status.empty()
                    st.success(f"Added {added} document(s). {duplicates} duplicate(s) skipped; {empty} empty/unreadable document(s).")
                    st.rerun()
                except Exception as e:
                    st.error(f"Processing failed: {e}")
        st.divider()
        st.caption("🔒 Documents and index are stored in a local SQLite database.")
        st.caption("The embedding model may download once on first launch.")

    brain = st.session_state.selected_brain
    stats = get_stats(brain)
    st.markdown(f"""
    <div class="hero">
      <div class="eyebrow">Local-first intelligence</div>
      <h1>Principia<span style="color:#a78bfa">.</span></h1>
      <p>Turn your documents into a searchable, connected knowledge base.</p>
    </div>
    """, unsafe_allow_html=True)

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Documents", stats["documents"])
    m2.metric("Knowledge chunks", stats["chunks"])
    m3.metric("Unique tags", stats["tags"])
    m4.metric("Workspace", brain)

    search_tab, chat_tab, library_tab, analytics_tab, settings_tab = st.tabs(
        ["⌕ Search", "✦ Ask Principia", "▤ Library", "◫ Insights", "⚙ Settings"]
    )

    with search_tab:
        st.subheader("Search your knowledge")
        st.caption("Find the passages that matter, with the original source attached.")
        with st.form("search_form"):
            query = st.text_input("Question or keywords", placeholder="What did the research say about customer retention?")
            c1, c2 = st.columns([2, 1])
            mode = c1.selectbox("Retrieval strategy", ["Hybrid", "Semantic", "Keyword"])
            top_k = c2.slider("Results", 1, 10, 5)
            submitted = st.form_submit_button("Search knowledge", type="primary", use_container_width=True)
        if submitted:
            if not query.strip():
                st.warning("Enter a question or keywords.")
            elif not stats["chunks"]:
                st.info("This workspace is empty. Add documents from the sidebar to get started.")
            else:
                try:
                    with st.spinner("Searching your local knowledge…"):
                        results = retrieve(query, brain, get_model(), mode, top_k)
                    st.session_state.last_search = {"query": query, "results": results, "brain": brain}
                except Exception as e:
                    st.error(f"Search failed: {e}")
        last = st.session_state.get("last_search", {})
        if last.get("brain") == brain:
            if last.get("results"):
                st.markdown(f"**{len(last['results'])} relevant passage(s)** for: {last['query']}")
                for i, result in enumerate(last["results"], 1):
                    with st.expander(f"{i}. {result['source']}  ·  relevance {result['score']:.3f}", expanded=(i == 1)):
                        st.caption(f"Added {result['added_at']} · Tags: {', '.join(json.loads(result['tags'])) or 'none'}")
                        st.write(result["text"])
            elif last:
                st.info("No matching passages found. Try broader terms or another retrieval strategy.")

    with chat_tab:
        st.subheader("Ask your documents")
        st.caption("Answers are grounded in retrieved passages. Open a citation to inspect the source.")
        history = st.session_state.chat_history.setdefault(brain, [])
        for msg in history:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])
        with st.form("chat_form", clear_on_submit=True):
            prompt = st.text_input("Ask a question", placeholder="Summarize the main findings…")
            send = st.form_submit_button("Ask", type="primary")
        if send:
            if not prompt.strip():
                st.warning("Type a question first.")
            elif not stats["chunks"]:
                st.info("Add documents to this workspace before asking questions.")
            else:
                history.append({"role": "user", "content": prompt.strip()})
                try:
                    with st.spinner("Finding evidence…"):
                        results = retrieve(prompt.strip(), brain, get_model(), "Hybrid", 5)
                    answer = answer_from_results(prompt.strip(), results)
                except Exception as e:
                    answer = f"Search failed safely: {e}"
                history.append({"role": "assistant", "content": answer})
                st.rerun()
        if history and st.button("Clear this chat"):
            st.session_state.chat_history[brain] = []
            st.rerun()

    with library_tab:
        st.subheader("Document library")
        if not stats["sources"]:
            st.info("Your library is empty. Upload a document in the sidebar.")
        else:
            df = pd.DataFrame(stats["sources"])
            df = df.rename(columns={"source": "Document", "added_at": "Added", "char_count": "Characters"})
            df["Tags"] = df["tags"].apply(lambda x: ", ".join(json.loads(x)) if x else "")
            st.dataframe(df[["Document", "Added", "Characters", "Tags"]], use_container_width=True, hide_index=True)
            st.caption("Duplicate files are detected by content hash within each workspace.")

    with analytics_tab:
        st.subheader("Knowledge insights")
        if not stats["sources"]:
            st.info("Add documents to see insights.")
        else:
            df = pd.DataFrame(stats["sources"])
            left, right = st.columns(2)
            with left:
                source_counts = df["source"].value_counts().reset_index()
                source_counts.columns = ["Document", "Chunks' source count"]
                fig = px.bar(source_counts, x="Document", y="Chunks' source count", title="Documents in this workspace")
                fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font_color="#e5e7eb")
                st.plotly_chart(fig, use_container_width=True)
            with right:
                df["Added"] = pd.to_datetime(df["added_at"], errors="coerce")
                timeline = df.dropna(subset=["Added"]).assign(Date=lambda x: x["Added"].dt.date).groupby("Date").size().reset_index(name="Documents")
                if not timeline.empty:
                    fig2 = px.line(timeline, x="Date", y="Documents", markers=True, title="Ingestion over time")
                    fig2.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font_color="#e5e7eb")
                    st.plotly_chart(fig2, use_container_width=True)
                else:
                    st.info("No valid ingestion timestamps yet.")
            all_tags = []
            for tag_json in df["tags"]:
                try:
                    all_tags.extend(json.loads(tag_json))
                except (ValueError, TypeError):
                    pass
            if all_tags:
                tag_counts = pd.Series(all_tags).value_counts().head(20).rename_axis("Tag").reset_index(name="Count")
                fig3 = px.bar(tag_counts, x="Count", y="Tag", orientation="h", title="Most common tags")
                fig3.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font_color="#e5e7eb")
                st.plotly_chart(fig3, use_container_width=True)

    with settings_tab:
        st.subheader("Workspace settings")
        st.markdown("#### Export")
        if st.button("Prepare JSON export"):
            with connect() as c:
                docs = c.execute("SELECT source, added_at, tags, char_count, content_hash FROM documents WHERE brain=? ORDER BY added_at", (brain,)).fetchall()
                chunks = c.execute("SELECT source, text FROM chunks JOIN documents ON documents.id=chunks.document_id WHERE chunks.brain=? ORDER BY documents.id, chunk_index", (brain,)).fetchall()
            payload = {
                "app": "Principia", "schema_version": 1, "workspace": brain,
                "exported_at": datetime.now().isoformat(timespec="seconds"),
                "documents": [dict(r) for r in docs],
                "chunks": [dict(r) for r in chunks],
            }
            st.download_button("Download workspace JSON", json.dumps(payload, indent=2), file_name=f"principia-{brain}.json", mime="application/json")
        st.divider()
        st.markdown("#### Data management")
        st.caption("Deleting a workspace permanently removes its documents and indexed chunks from the local database.")
        if brain == "default":
            st.info("The default workspace is protected from deletion.")
        else:
            confirm = st.checkbox(f"I understand that deleting '{brain}' is permanent.", key=f"confirm_delete_{brain}")
            if st.button("Delete this workspace", disabled=not confirm):
                with connect() as c:
                    c.execute("DELETE FROM brains WHERE name=?", (brain,))
                st.session_state.selected_brain = "default"
                st.session_state.chat_history.pop(brain, None)
                st.success("Workspace deleted.")
                st.rerun()
        st.divider()
        st.markdown("#### Product notes")
        st.write("- Local SQLite persistence for document metadata, chunks, and embeddings.")
        st.write("- Hybrid retrieval combines semantic similarity with keyword matching.")
        st.write("- Answers are extractive and source-linked; no cloud LLM is called.")
        st.write("- Scanned PDFs need OCR before their text can be indexed.")

if __name__ == "__main__":
    main()
