"""Principia — a local-first personal knowledge workspace.
Run: pip install -r requirements.txt
Then: streamlit run principia.py
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from datetime import datetime, date
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ============================================================
# Configuration
# ============================================================

APP_NAME = "Principia"
MODEL_NAME = "all-MiniLM-L6-v2"
CHUNK_SIZE = 240
OVERLAP = 40
MAX_MB = 40

# Use a writable location rather than the source-code directory.
# Note: the home directory on Streamlit Cloud may not persist
# across rebuilds. Use a persistent database service for production.
DATA_DIR = Path.home() / ".principia"
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / "principia.sqlite3"

st.set_page_config(
    page_title="Principia — Think clearly",
    page_icon="🧠",
    layout="wide",
)

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap');
.stApp {
    background: radial-gradient(ellipse at 10% 0%, #20274b 0%, #11172b 45%, #090e1b 100%);
    color: #f3f5ff;
    font-family: 'DM Sans', sans-serif;
}
[data-testid="stSidebar"] {
    background: rgba(7,11,23,.94);
    border-right: 1px solid #ffffff18;
}
.block-container { max-width: 1480px; padding-top: 1.3rem; }
h1,h2,h3 { font-family: 'Space Grotesk', sans-serif; letter-spacing: -.035em; }
.hero {
    padding: 1.8rem 2rem;
    border: 1px solid #a99aff44;
    border-radius: 24px;
    background: linear-gradient(115deg,#4f46e538,#8b5cf61c,#ffffff06);
    margin-bottom: 1.2rem;
}
.hero small {
    color: #c9c0ff;
    letter-spacing: .16em;
    text-transform: uppercase;
    font-weight: 700;
}
.hero h1 {
    font-size: clamp(2.4rem,4vw,3.4rem);
    margin: .35rem 0;
    color: white;
}
.hero p { color: #c5cce0; font-size: 1rem; margin: 0; max-width: 800px; }
div[data-testid="stMetric"] {
    background: #ffffff08;
    border: 1px solid #ffffff15;
    padding: 14px;
    border-radius: 14px;
}
.stButton button,.stDownloadButton button {
    border-radius: 10px;
    border: 1px solid #a99aff44;
    font-weight: 650;
}
.stButton button[kind="primary"] {
    background: linear-gradient(100deg,#5b4be5,#8b5cf6);
    color: white;
    border: 0;
}
.stTextInput input,.stTextArea textarea {
    background: #ffffff0b;
    color: white;
    border-color: #ffffff20;
    border-radius: 10px;
}
div[data-testid="stExpander"] { border: 1px solid #ffffff18; border-radius: 12px; }
</style>
""",
    unsafe_allow_html=True,
)


# ============================================================
# Database and safe schema migration
# ============================================================

def connect():
    """Open a SQLite connection with consistent settings."""
    connection = sqlite3.connect(str(DB_PATH), timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def table_columns(connection, table_name):
    """Return existing column names for a table."""
    return {
        row["name"]
        for row in connection.execute(
            f"PRAGMA table_info({table_name})"
        ).fetchall()
    }


def ensure_column(connection, table, column, definition):
    """Add a missing column to an existing table."""
    existing = table_columns(connection, table)
    if column not in existing:
        connection.execute(
            f"ALTER TABLE {table} ADD COLUMN {column} {definition}"
        )


def init_db():
    """Create tables, migrate older databases, then create indexes."""
    with connect() as connection:
        # Create tables first. Do not create indexes here: older versions
        # may have tables whose columns differ from the current schema.
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS workspaces (
                name TEXT PRIMARY KEY,
                created_at TEXT NOT NULL DEFAULT '',
                purpose TEXT NOT NULL DEFAULT 'General'
            );

            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY,
                workspace TEXT NOT NULL DEFAULT 'My Knowledge',
                title TEXT NOT NULL DEFAULT 'Untitled',
                digest TEXT NOT NULL DEFAULT '',
                added_at TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '[]',
                chars INTEGER NOT NULL DEFAULT 0,
                pages INTEGER,
                UNIQUE(workspace, digest)
            );

            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY,
                document_id INTEGER NOT NULL DEFAULT 0,
                workspace TEXT NOT NULL DEFAULT 'My Knowledge',
                part INTEGER NOT NULL DEFAULT 0,
                content TEXT NOT NULL DEFAULT '',
                embedding BLOB,
                dim INTEGER
            );

            CREATE TABLE IF NOT EXISTS notes (
                id INTEGER PRIMARY KEY,
                workspace TEXT NOT NULL DEFAULT 'My Knowledge',
                title TEXT NOT NULL DEFAULT 'Untitled',
                content TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY,
                workspace TEXT NOT NULL DEFAULT 'My Knowledge',
                title TEXT NOT NULL DEFAULT 'Untitled',
                details TEXT NOT NULL DEFAULT '',
                due TEXT,
                priority TEXT NOT NULL DEFAULT 'Normal',
                status TEXT NOT NULL DEFAULT 'Open',
                created_at TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS activity (
                id INTEGER PRIMARY KEY,
                workspace TEXT NOT NULL DEFAULT 'My Knowledge',
                action TEXT NOT NULL DEFAULT '',
                detail TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT ''
            );
            """
        )

        # Migrate older versions. CREATE TABLE IF NOT EXISTS does not
        # modify a table that already exists.
        migrations = {
            "workspaces": {
                "created_at": "TEXT NOT NULL DEFAULT ''",
                "purpose": "TEXT NOT NULL DEFAULT 'General'",
            },
            "documents": {
                "workspace": "TEXT NOT NULL DEFAULT 'My Knowledge'",
                "title": "TEXT NOT NULL DEFAULT 'Untitled'",
                "digest": "TEXT NOT NULL DEFAULT ''",
                "added_at": "TEXT NOT NULL DEFAULT ''",
                "tags": "TEXT NOT NULL DEFAULT '[]'",
                "chars": "INTEGER NOT NULL DEFAULT 0",
                "pages": "INTEGER",
            },
            "chunks": {
                "document_id": "INTEGER NOT NULL DEFAULT 0",
                "workspace": "TEXT NOT NULL DEFAULT 'My Knowledge'",
                "part": "INTEGER NOT NULL DEFAULT 0",
                "content": "TEXT NOT NULL DEFAULT ''",
                "embedding": "BLOB",
                "dim": "INTEGER",
            },
            "notes": {
                "workspace": "TEXT NOT NULL DEFAULT 'My Knowledge'",
                "title": "TEXT NOT NULL DEFAULT 'Untitled'",
                "content": "TEXT NOT NULL DEFAULT ''",
                "tags": "TEXT NOT NULL DEFAULT '[]'",
                "created_at": "TEXT NOT NULL DEFAULT ''",
                "updated_at": "TEXT NOT NULL DEFAULT ''",
            },
            "tasks": {
                "workspace": "TEXT NOT NULL DEFAULT 'My Knowledge'",
                "title": "TEXT NOT NULL DEFAULT 'Untitled'",
                "details": "TEXT NOT NULL DEFAULT ''",
                "due": "TEXT",
                "priority": "TEXT NOT NULL DEFAULT 'Normal'",
                "status": "TEXT NOT NULL DEFAULT 'Open'",
                "created_at": "TEXT NOT NULL DEFAULT ''",
            },
            "activity": {
                "workspace": "TEXT NOT NULL DEFAULT 'My Knowledge'",
                "action": "TEXT NOT NULL DEFAULT ''",
                "detail": "TEXT NOT NULL DEFAULT ''",
                "created_at": "TEXT NOT NULL DEFAULT ''",
            },
        }

        for table, columns in migrations.items():
            for column, definition in columns.items():
                ensure_column(connection, table, column, definition)

        # Ensure the default workspace exists before normal app usage.
        connection.execute(
            """
            INSERT OR IGNORE INTO workspaces(name, created_at, purpose)
            VALUES (?, ?, ?)
            """,
            (
                "My Knowledge",
                datetime.now().isoformat(timespec="seconds"),
                "Personal",
            ),
        )

        # Repair legacy rows with blank workspace values.
        for table in ("documents", "chunks", "notes", "tasks", "activity"):
            connection.execute(
                f"""
                UPDATE {table}
                SET workspace = 'My Knowledge'
                WHERE workspace IS NULL OR TRIM(workspace) = ''
                """
            )

        # Index creation happens only after the migration.
        connection.execute(
            "CREATE INDEX IF NOT EXISTS chunks_ws ON chunks(workspace)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS documents_ws ON documents(workspace)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS notes_ws ON notes(workspace)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS tasks_ws ON tasks(workspace)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS activity_ws ON activity(workspace)"
        )


def log(ws, action, detail):
    with connect() as connection:
        connection.execute(
            """
            INSERT INTO activity(workspace, action, detail, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (
                ws,
                action,
                detail,
                datetime.now().isoformat(timespec="seconds"),
            ),
        )


# ============================================================
# Utility functions
# ============================================================

def clean_name(value):
    value = re.sub(r"[^a-zA-Z0-9 _-]", "", value).strip()
    value = re.sub(r"\s+", "-", value)[:48]
    if not value:
        raise ValueError(
            "Use letters, numbers, spaces, hyphens, or underscores."
        )
    return value


@st.cache_resource(show_spinner="Loading local embedding model…")
def model():
    return SentenceTransformer(MODEL_NAME)


def parse_file(upload):
    raw = upload.getvalue()
    name = upload.name.lower()

    if len(raw) > MAX_MB * 1024 * 1024:
        raise ValueError(f"File exceeds {MAX_MB} MB.")

    if name.endswith(".pdf"):
        pdf = PdfReader(io.BytesIO(raw), strict=False)
        text = "\n\n".join(
            page.extract_text() or "" for page in pdf.pages
        ).strip()
        return text, len(pdf.pages)

    if name.endswith(".docx"):
        document = Document(io.BytesIO(raw))
        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )
        return "\n\n".join(parts).strip(), None

    return raw.decode("utf-8-sig", errors="replace").strip(), None


def chunk_text(text):
    words = re.findall(r"\S+", text)
    chunks = []
    start = 0

    while start < len(words):
        end = min(start + CHUNK_SIZE, len(words))
        chunk = " ".join(words[start:end]).strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(words):
            break

        start = max(end - OVERLAP, start + 1)

    return chunks


def ingest(ws, upload, tags, encoder):
    text, pages = parse_file(upload)

    if not text:
        return "empty", 0

    chunks = chunk_text(text)
    digest = hashlib.sha256(upload.getvalue()).hexdigest()

    with connect() as connection:
        duplicate = connection.execute(
            "SELECT 1 FROM documents WHERE workspace=? AND digest=?",
            (ws, digest),
        ).fetchone()

        if duplicate:
            return "duplicate", 0

        cursor = connection.execute(
            """
            INSERT INTO documents
            (workspace, title, digest, added_at, tags, chars, pages)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ws,
                upload.name,
                digest,
                datetime.now().isoformat(timespec="seconds"),
                json.dumps(tags),
                len(text),
                pages,
            ),
        )
        document_id = cursor.lastrowid

    vectors = encoder.encode(
        chunks,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    with connect() as connection:
        connection.executemany(
            """
            INSERT INTO chunks
            (document_id, workspace, part, content, embedding, dim)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    document_id,
                    ws,
                    index,
                    chunk,
                    np.asarray(vector, dtype=np.float32).tobytes(),
                    len(vector),
                )
                for index, (chunk, vector) in enumerate(zip(chunks, vectors))
            ],
        )

    log(ws, "Document indexed", upload.name)
    return "added", len(chunks)


def workspace_rows(ws):
    with connect() as connection:
        docs = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM documents WHERE workspace=? ORDER BY added_at DESC",
                (ws,),
            )
        ]

        chunks = [
            dict(row)
            for row in connection.execute(
                """
                SELECT ch.*, d.title, d.added_at AS doc_date, d.tags, d.pages
                FROM chunks ch
                JOIN documents d ON d.id = ch.document_id
                WHERE ch.workspace=?
                ORDER BY ch.id
                """,
                (ws,),
            )
        ]

        notes = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM notes WHERE workspace=? ORDER BY updated_at DESC",
                (ws,),
            )
        ]

        tasks = [
            dict(row)
            for row in connection.execute(
                """
                SELECT * FROM tasks
                WHERE workspace=?
                ORDER BY CASE priority
                    WHEN 'High' THEN 0
                    WHEN 'Normal' THEN 1
                    ELSE 2
                END, due
                """,
                (ws,),
            )
        ]

        activity = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM activity WHERE workspace=? ORDER BY id DESC LIMIT 20",
                (ws,),
            )
        ]

    return docs, chunks, notes, tasks, activity


def search(query, ws, encoder, mode="Hybrid", limit=5, tag="All tags"):
    _, rows, _, _, _ = workspace_rows(ws)

    if tag != "All tags":
        filtered = []
        for row in rows:
            try:
                row_tags = json.loads(row["tags"] or "[]")
            except (ValueError, TypeError):
                row_tags = []
            if tag in row_tags:
                filtered.append(row)
        rows = filtered

    if not query.strip() or not rows:
        return []

    scores = np.zeros(len(rows), dtype=float)

    if mode in ("Hybrid", "Semantic"):
        query_vector = np.asarray(
            encoder.encode([query], normalize_embeddings=True)[0],
            dtype=np.float32,
        )

        for index, row in enumerate(rows):
            blob = row.get("embedding")
            dimension = row.get("dim")

            if blob and dimension == len(query_vector):
                vector = np.frombuffer(blob, dtype=np.float32)
                if len(vector) == len(query_vector) and np.isfinite(vector).all():
                    scores[index] = float(np.dot(query_vector, vector))

        if mode == "Hybrid":
            scores *= 0.72

    if mode in ("Hybrid", "Keyword"):
        try:
            vectorizer = TfidfVectorizer(
                ngram_range=(1, 2),
                stop_words="english",
                sublinear_tf=True,
            )
            matrix = vectorizer.fit_transform(
                [row["content"] for row in rows]
            )
            keyword_scores = cosine_similarity(
                vectorizer.transform([query]), matrix
            )[0]
        except ValueError:
            keyword_scores = np.zeros(len(rows))

        scores = (
            keyword_scores
            if mode == "Keyword"
            else scores + 0.28 * keyword_scores
        )

    output = []
    for index in np.argsort(scores)[::-1][:limit]:
        if scores[index] > 0:
            item = rows[int(index)].copy()
            item["score"] = float(scores[index])
            output.append(item)

    return output


def evidence_answer(question, results):
    if not results:
        return (
            "I couldn't find supporting passages. Try different wording, "
            "remove the tag filter, or add a relevant document."
        )

    terms = set(re.findall(r"\b[a-zA-Z0-9]{3,}\b", question.lower()))
    pool = []

    for row in results[:5]:
        for sentence in re.split(r"(?<=[.!?])\s+", row["content"]):
            sentence = sentence.strip()
            if len(sentence) < 24:
                continue

            sentence_terms = set(
                re.findall(r"\b[a-zA-Z0-9]{3,}\b", sentence.lower())
            )
            score = (
                len(terms & sentence_terms) / max(1, len(terms))
                + min(len(sentence), 400) / 4000
            )
            pool.append((score, sentence, row))

    pool.sort(key=lambda item: item[0], reverse=True)

    chosen = []
    seen = set()

    for _, sentence, row in pool:
        if sentence.lower() not in seen:
            chosen.append((sentence, row))
            seen.add(sentence.lower())
        if len(chosen) == 3:
            break

    if not chosen:
        chosen = [(row["content"][:400], row) for row in results[:2]]

    lines = ["**Relevant evidence from your workspace**"]

    for sentence, row in chosen:
        lines.append(
            f"- “{sentence}”\n\n"
            f"  **Source:** {row['title']} · "
            f"relevance `{row['score']:.3f}`"
        )

    lines.append(
        "*These are extracted passages, not fabricated answers from a generative AI model.*"
    )
    return "\n\n".join(lines)


def export_json(ws):
    docs, chunks, notes, tasks, _ = workspace_rows(ws)

    for document in docs:
        document.pop("id", None)

    clean_chunks = [
        {
            "document": row["title"],
            "part": row["part"],
            "content": row["content"],
        }
        for row in chunks
    ]

    clean_notes = [
        {
            key: note[key]
            for key in (
                "title",
                "content",
                "tags",
                "created_at",
                "updated_at",
            )
        }
        for note in notes
    ]

    clean_tasks = [
        {
            key: task[key]
            for key in (
                "title",
                "details",
                "due",
                "priority",
                "status",
                "created_at",
            )
        }
        for task in tasks
    ]

    payload = {
        "product": "Principia",
        "schema_version": 1,
        "workspace": ws,
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "documents": docs,
        "chunks": clean_chunks,
        "notes": clean_notes,
        "tasks": clean_tasks,
    }

    return json.dumps(
        payload, ensure_ascii=False, indent=2
    ).encode("utf-8")


# ============================================================
# Streamlit interface
# ============================================================

def main():
    try:
        init_db()
    except sqlite3.Error as exc:
        st.error("Principia could not initialize its database.")
        st.code(str(exc))
        st.info(
            "Check the database path and schema. Do not delete your existing "
            "database unless you have a backup."
        )
        st.stop()

    for key, default in (
        ("ws", "My Knowledge"),
        ("chats", {}),
        ("last", {}),
    ):
        if key not in st.session_state:
            st.session_state[key] = default

    with st.sidebar:
        st.markdown("## 🧠 Principia")
        st.caption("Ideas → knowledge → action")

        with connect() as connection:
            names = [
                row["name"]
                for row in connection.execute(
                    "SELECT name FROM workspaces ORDER BY name"
                )
            ]

        if not names:
            names = ["My Knowledge"]

        if st.session_state.ws not in names:
            st.session_state.ws = names[0]

        ws = st.selectbox(
            "Workspace",
            names,
            index=names.index(st.session_state.ws),
        )
        st.session_state.ws = ws

        with st.expander("＋ New workspace"):
            workspace_name = st.text_input(
                "Name",
                placeholder="School, Team, Personal",
            )
            purpose = st.selectbox(
                "Purpose",
                ["Personal", "School", "Family", "Work", "Research", "General"],
            )

            if st.button("Create workspace", width="stretch"):
                try:
                    name = clean_name(workspace_name)

                    with connect() as connection:
                        connection.execute(
                            """
                            INSERT INTO workspaces(name, created_at, purpose)
                            VALUES (?, ?, ?)
                            """,
                            (
                                name,
                                datetime.now().isoformat(timespec="seconds"),
                                purpose,
                            ),
                        )

                    log(name, "Workspace created", purpose)
                    st.session_state.ws = name
                    st.rerun()

                except (ValueError, sqlite3.IntegrityError) as exc:
                    st.error(str(exc))

        st.divider()
        st.markdown("### Add knowledge")

        files = st.file_uploader(
            "PDF, DOCX, TXT, Markdown",
            type=["pdf", "docx", "txt", "md"],
            accept_multiple_files=True,
        )

        tags_raw = st.text_input(
            "Tags",
            placeholder="study, project, reference",
        )
        tags = list(
            dict.fromkeys(
                item.strip()
                for item in tags_raw.split(",")
                if item.strip()
            )
        )

        if st.button(
            "Index documents",
            type="primary",
            width="stretch",
        ):
            if not files:
                st.warning("Choose at least one file.")
            else:
                try:
                    encoder = model()
                    progress = st.progress(0)
                    added = duplicates = empty = failed = 0

                    for index, uploaded_file in enumerate(files):
                        try:
                            result, _ = ingest(
                                ws, uploaded_file, tags, encoder
                            )
                            if result == "added":
                                added += 1
                            elif result == "duplicate":
                                duplicates += 1
                            else:
                                empty += 1
                        except Exception as exc:
                            failed += 1
                            st.warning(f"{uploaded_file.name}: {exc}")

                        progress.progress((index + 1) / len(files))

                    st.success(
                        f"Added {added}; duplicates {duplicates}; "
                        f"empty {empty}; failed {failed}."
                    )
                    st.rerun()

                except Exception as exc:
                    st.error(f"Indexing failed: {exc}")

        st.caption(
            "SQLite storage. The embedding model may download on first use."
        )
        st.divider()
        st.markdown("### Quick capture")

        with st.form("capture", clear_on_submit=True):
            note_title = st.text_input(
                "Note title",
                placeholder="An idea worth keeping",
            )
            note_content = st.text_area(
                "Note",
                placeholder="Capture a thought, decision, lesson…",
                height=90,
            )
            note_tags = st.text_input(
                "Tags",
                placeholder="idea, learning",
            )
            save_note = st.form_submit_button(
                "Save note",
                width="stretch",
            )

        if save_note:
            if not note_title.strip() or not note_content.strip():
                st.warning("Add a title and note content.")
            else:
                now = datetime.now().isoformat(timespec="seconds")
                parsed_tags = [
                    item.strip()
                    for item in note_tags.split(",")
                    if item.strip()
                ]

                with connect() as connection:
                    connection.execute(
                        """
                        INSERT INTO notes
                        (workspace, title, content, tags, created_at, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            ws,
                            note_title.strip(),
                            note_content.strip(),
                            json.dumps(parsed_tags),
                            now,
                            now,
                        ),
                    )

                log(ws, "Note captured", note_title.strip())
                st.success("Saved.")
                st.rerun()

    docs, chunks, notes, tasks, activity = workspace_rows(ws)

    all_tags = set()
    for row in docs + notes:
        try:
            all_tags.update(json.loads(row["tags"] or "[]"))
        except (ValueError, TypeError):
            pass

    st.markdown(
        f"""
<div class="hero">
  <small>PERSONAL KNOWLEDGE OS · {ws}</small>
  <h1>Think clearly. <span style="color:#a99aff">Find anything.</span></h1>
  <p>One calm workspace for learning, family life, research, and meaningful work.
  Keep the evidence close and the next action clear.</p>
</div>
""",
        unsafe_allow_html=True,
    )

    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("Documents", len(docs))
    col2.metric("Indexed passages", len(chunks))
    col3.metric("Notes", len(notes))
    col4.metric("Open tasks", sum(task["status"] != "Done" for task in tasks))
    col5.metric("Tags", len(all_tags))

    home, ask, library, notes_tab, tasks_tab, insights, settings = st.tabs(
        [
            "⌂ Overview",
            "✦ Ask & Search",
            "▤ Library",
            "✎ Notes",
            "✓ Tasks",
            "◫ Insights",
            "⚙ Settings",
        ]
    )

    # Overview
    with home:
        left, right = st.columns([1.2, 1])

        with left:
            st.subheader("Your next move")
            open_tasks = [
                task for task in tasks if task["status"] != "Done"
            ]

            if open_tasks:
                for task in open_tasks[:5]:
                    icon = (
                        "🔴" if task["priority"] == "High"
                        else "🟡" if task["priority"] == "Normal"
                        else "⚪"
                    )
                    st.markdown(
                        f"{icon} **{task['title']}** · {task['status']}"
                    )
                    st.caption(
                        f"Due: {task['due'] or 'not set'} · "
                        f"{task['priority']}"
                    )
            else:
                st.info(
                    "No open tasks. Add a small, specific next action in Tasks."
                )

            st.markdown("#### Suggested ways to use Principia")
            st.markdown(
                "- **Students:** upload readings, then search for concepts and examples.\n"
                "- **Families:** organize shared reference material and checklists.\n"
                "- **Professionals:** collect briefs, meeting notes, research, and next actions."
            )

        with right:
            st.subheader("Recently added")
            if not docs and not notes:
                st.info(
                    "Your workspace is ready. Upload a file or save a note "
                    "from the sidebar."
                )

            for document in docs[:4]:
                st.markdown(f"📄 **{document['title']}**")
                st.caption(
                    f"{document['added_at'][:16]} · "
                    f"{document['chars']:,} characters"
                )

            for note in notes[:3]:
                st.markdown(f"📝 **{note['title']}**")
                st.caption(f"Updated {note['updated_at'][:16]}")

        st.divider()
        st.subheader("Recent activity")

        if activity:
            st.dataframe(
                pd.DataFrame(activity)[["created_at", "action", "detail"]],
                width="stretch",
                hide_index=True,
            )
        else:
            st.caption("Your activity appears here as you work.")

    # Search and chat
    with ask:
        st.subheader("Ask your knowledge")
        st.caption(
            "Search by meaning or exact wording. Review the original passage "
            "before relying on it."
        )

        tag_options = ["All tags"] + sorted(all_tags)
        selected_tag = st.selectbox(
            "Filter by tag",
            tag_options,
            key=f"tag_{ws}",
        )

        with st.form("searchform"):
            query = st.text_input(
                "What are you trying to find?",
                placeholder="What are the key ideas in my notes about energy?",
            )
            mode = st.selectbox(
                "Search strategy",
                ["Hybrid", "Semantic", "Keyword"],
            )
            limit = st.slider("Passages", 1, 10, 5)
            search_clicked = st.form_submit_button(
                "Search workspace",
                type="primary",
                width="stretch",
            )

        if search_clicked:
            if not query.strip():
                st.warning("Enter a question or keywords.")
            elif not chunks:
                st.info("Add a document first.")
            else:
                try:
                    with st.spinner("Searching locally…"):
                        found = search(
                            query,
                            ws,
                            model(),
                            mode,
                            limit,
                            selected_tag,
                        )
                    st.session_state.last[ws] = {
                        "q": query,
                        "mode": mode,
                        "found": found,
                    }
                except Exception as exc:
                    st.error(f"Search failed: {exc}")

        last = st.session_state.last.get(ws)

        if last:
            st.markdown(
                f"#### {len(last['found'])} passage(s) · {last['mode']}"
            )

            for index, result in enumerate(last["found"], 1):
                with st.expander(
                    f"{index}. {result['title']} · "
                    f"relevance {result['score']:.3f}",
                    expanded=index == 1,
                ):
                    st.caption(
                        f"Added {result['doc_date'][:16]} · tags: "
                        f"{', '.join(json.loads(result['tags'] or '[]')) or 'none'}"
                    )
                    st.write(result["content"])

            st.markdown("#### Evidence summary")
            st.markdown(evidence_answer(last["q"], last["found"]))

        st.divider()
        st.subheader("Chat with your library")

        history = st.session_state.chats.setdefault(ws, [])

        for message in history:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        with st.form("chatform", clear_on_submit=True):
            prompt = st.text_input(
                "Ask a question",
                placeholder="Compare the main ideas across my documents…",
            )
            send = st.form_submit_button(
                "Ask Principia",
                type="primary",
            )

        if send:
            if not prompt.strip():
                st.warning("Type a question first.")
            elif not chunks:
                st.info("Add documents first.")
            else:
                history.append(
                    {"role": "user", "content": prompt.strip()}
                )
                try:
                    response = evidence_answer(
                        prompt,
                        search(prompt, ws, model(), "Hybrid", 5),
                    )
                except Exception as exc:
                    response = f"Search failed: {exc}"

                history.append(
                    {"role": "assistant", "content": response}
                )
                st.rerun()

        if history and st.button("Clear chat for this workspace"):
            st.session_state.chats[ws] = []
            st.rerun()

    # Document library
    with library:
        st.subheader("Document library")

        if not docs:
            st.info(
                "No documents yet. Upload PDF, Word, TXT, or Markdown files "
                "in the sidebar."
            )
        else:
            view = [
                {
                    "Title": document["title"],
                    "Added": document["added_at"],
                    "Characters": document["chars"],
                    "Pages": document["pages"] or "—",
                    "Tags": ", ".join(
                        json.loads(document["tags"] or "[]")
                    ),
                }
                for document in docs
            ]

            st.dataframe(
                pd.DataFrame(view),
                width="stretch",
                hide_index=True,
            )
            st.caption(
                "Identical file content is only indexed once per workspace."
            )

            with st.expander("Delete a document and its indexed passages"):
                document_titles = list(
                    dict.fromkeys(document["title"] for document in docs)
                )
                selected_title = st.selectbox(
                    "Document title",
                    document_titles,
                )
                confirm_delete = st.checkbox(
                    "I understand this also deletes its search passages.",
                    key=f"del_doc_{ws}_{selected_title}",
                )

                if st.button(
                    "Delete document",
                    disabled=not confirm_delete,
                ):
                    with connect() as connection:
                        connection.execute(
                            """
                            DELETE FROM chunks
                            WHERE document_id IN (
                                SELECT id FROM documents
                                WHERE workspace=? AND title=?
                            )
                            """,
                            (ws, selected_title),
                        )
                        connection.execute(
                            "DELETE FROM documents WHERE workspace=? AND title=?",
                            (ws, selected_title),
                        )

                    log(ws, "Document deleted", selected_title)
                    st.rerun()

    # Notes
    with notes_tab:
        st.subheader("Notes and ideas")

        if not notes:
            st.info("Use Quick capture in the sidebar to save your first note.")

        for note in notes:
            with st.expander(
                f"{note['title']} · {note['updated_at'][:16]}"
            ):
                with st.form(f"editnote_{note['id']}"):
                    edit_title = st.text_input("Title", note["title"])
                    edit_content = st.text_area(
                        "Content",
                        note["content"],
                        height=130,
                    )
                    edit_tags = st.text_input(
                        "Tags",
                        ", ".join(json.loads(note["tags"] or "[]")),
                    )
                    update_note = st.form_submit_button("Save changes")

                if update_note:
                    parsed_tags = [
                        item.strip()
                        for item in edit_tags.split(",")
                        if item.strip()
                    ]

                    with connect() as connection:
                        connection.execute(
                            """
                            UPDATE notes
                            SET title=?, content=?, tags=?, updated_at=?
                            WHERE id=? AND workspace=?
                            """,
                            (
                                edit_title.strip(),
                                edit_content.strip(),
                                json.dumps(parsed_tags),
                                datetime.now().isoformat(timespec="seconds"),
                                note["id"],
                                ws,
                            ),
                        )

                    log(ws, "Note edited", edit_title.strip())
                    st.rerun()

                if st.button("Delete note", key=f"del_note_{note['id']}"):
                    with connect() as connection:
                        connection.execute(
                            "DELETE FROM notes WHERE id=? AND workspace=?",
                            (note["id"], ws),
                        )
                    log(ws, "Note deleted", note["title"])
                    st.rerun()

    # Tasks
    with tasks_tab:
        st.subheader("Tasks and next actions")

        with st.form("newtask", clear_on_submit=True):
            task_title = st.text_input(
                "Task",
                placeholder="Review chapter 3",
            )
            task_details = st.text_area("Details (optional)", height=70)

            col_due, col_priority = st.columns(2)
            due = col_due.date_input(
                "Due date",
                value=None,
                format="YYYY-MM-DD",
            )
            priority = col_priority.selectbox(
                "Priority",
                ["Normal", "High", "Low"],
            )
            add_task = st.form_submit_button(
                "Add task",
                type="primary",
                width="stretch",
            )

        if add_task:
            if not task_title.strip():
                st.warning("Give the task a title.")
            else:
                due_value = (
                    due.isoformat() if isinstance(due, date) else None
                )

                with connect() as connection:
                    connection.execute(
                        """
                        INSERT INTO tasks
                        (workspace, title, details, due, priority, status, created_at)
                        VALUES (?, ?, ?, ?, ?, 'Open', ?)
                        """,
                        (
                            ws,
                            task_title.strip(),
                            task_details.strip(),
                            due_value,
                            priority,
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )

                log(ws, "Task added", task_title.strip())
                st.rerun()

        if not tasks:
            st.info("No tasks yet.")

        for task in tasks:
            with st.container(border=True):
                col_title, col_status, col_delete = st.columns([4, 1.4, 1])

                col_title.markdown(f"**{task['title']}**")
                col_title.caption(
                    f"{task['details'] or 'No details'} · "
                    f"{task['priority']} priority · "
                    f"due {task['due'] or 'not set'}"
                )

                statuses = ["Open", "In progress", "Done"]
                current_status = (
                    task["status"]
                    if task["status"] in statuses
                    else "Open"
                )

                selected_status = col_status.selectbox(
                    "Status",
                    statuses,
                    index=statuses.index(current_status),
                    key=f"status_{task['id']}",
                    label_visibility="collapsed",
                )

                if selected_status != current_status:
                    with connect() as connection:
                        connection.execute(
                            """
                            UPDATE tasks SET status=?
                            WHERE id=? AND workspace=?
                            """,
                            (selected_status, task["id"], ws),
                        )

                    log(
                        ws,
                        "Task updated",
                        f"{task['title']}: {selected_status}",
                    )
                    st.rerun()

                if col_delete.button(
                    "Delete",
                    key=f"taskdel_{task['id']}",
                ):
                    with connect() as connection:
                        connection.execute(
                            "DELETE FROM tasks WHERE id=? AND workspace=?",
                            (task["id"], ws),
                        )

                    log(ws, "Task deleted", task["title"])
                    st.rerun()

    # Insights
    with insights:
        st.subheader("Knowledge insights")

        if not docs and not notes:
            st.info("Add documents or notes to see insights.")
        else:
            col_left, col_right = st.columns(2)

            if docs:
                dataframe = pd.DataFrame(docs)

                with col_left:
                    document_counts = (
                        dataframe.groupby("title", as_index=False)
                        .size()
                    )
                    figure = px.bar(
                        document_counts,
                        x="title",
                        y="size",
                        title="Documents in workspace",
                        labels={"title": "Document", "size": "Count"},
                    )
                    figure.update_layout(
                        paper_bgcolor="rgba(0,0,0,0)",
                        plot_bgcolor="rgba(0,0,0,0)",
                        font_color="#e8eaff",
                    )
                    st.plotly_chart(figure, width="stretch")

                with col_right:
                    dataframe["day"] = pd.to_datetime(
                        dataframe["added_at"],
                        errors="coerce",
                    ).dt.date

                    timeline = (
                        dataframe.dropna(subset=["day"])
                        .groupby("day")
                        .size()
                        .reset_index(name="Documents")
                    )

                    if not timeline.empty:
                        figure = px.line(
                            timeline,
                            x="day",
                            y="Documents",
                            markers=True,
                            title="Knowledge added over time",
                        )
                        figure.update_layout(
                            paper_bgcolor="rgba(0,0,0,0)",
                            plot_bgcolor="rgba(0,0,0,0)",
                            font_color="#e8eaff",
                        )
                        st.plotly_chart(figure, width="stretch")

            all_tag_values = []
            for row in docs + notes:
                try:
                    all_tag_values.extend(json.loads(row["tags"] or "[]"))
                except (ValueError, TypeError):
                    pass

            if all_tag_values:
                tag_counts = (
                    pd.Series(all_tag_values)
                    .value_counts()
                    .head(15)
                    .rename_axis("Tag")
                    .reset_index(name="Count")
                )

                figure = px.bar(
                    tag_counts,
                    x="Count",
                    y="Tag",
                    orientation="h",
                    title="Most-used tags",
                )
                figure.update_layout(
                    paper_bgcolor="rgba(0,0,0,0)",
                    plot_bgcolor="rgba(0,0,0,0)",
                    font_color="#e8eaff",
                )
                st.plotly_chart(figure, width="stretch")

    # Settings and exports
    with settings:
        st.subheader("Portability and workspace settings")

        safe_workspace = re.sub(r"[^a-zA-Z0-9_-]", "-", ws)

        st.download_button(
            "Export workspace JSON",
            data=export_json(ws),
            file_name=f"principia-{safe_workspace}.json",
            mime="application/json",
            width="stretch",
        )

        if DB_PATH.exists():
            try:
                st.download_button(
                    "Download SQLite backup",
                    data=DB_PATH.read_bytes(),
                    file_name="principia-backup.sqlite3",
                    mime="application/octet-stream",
                    width="stretch",
                )
            except OSError as exc:
                st.warning(f"Could not read the database backup: {exc}")

        st.divider()

        if ws == "My Knowledge":
            st.info("The default workspace is protected from deletion.")
        else:
            confirm_workspace_delete = st.checkbox(
                f"Permanently delete {ws} and its documents, notes, and tasks.",
                key=f"delete_ws_{ws}",
            )

            if st.button(
                "Delete workspace",
                disabled=not confirm_workspace_delete,
            ):
                with connect() as connection:
                    # Delete related records explicitly for legacy databases
                    # that may not have foreign-key constraints.
                    document_ids = [
                        row["id"]
                        for row in connection.execute(
                            "SELECT id FROM documents WHERE workspace=?",
                            (ws,),
                        )
                    ]

                    if document_ids:
                        placeholders = ",".join("?" for _ in document_ids)
                        connection.execute(
                            f"DELETE FROM chunks WHERE document_id IN ({placeholders})",
                            document_ids,
                        )

                    for table in ("documents", "notes", "tasks", "activity"):
                        connection.execute(
                            f"DELETE FROM {table} WHERE workspace=?",
                            (ws,),
                        )

                    connection.execute(
                        "DELETE FROM workspaces WHERE name=?",
                        (ws,),
                    )

                st.session_state.chats.pop(ws, None)
                st.session_state.last.pop(ws, None)
                st.session_state.ws = "My Knowledge"
                st.rerun()

        st.markdown("#### Capability and launch notes")
        st.markdown(
            "- SQLite stores workspace metadata, extracted text, chunks, and embeddings.\n"
            "- Semantic search uses `all-MiniLM-L6-v2`; the model may download on first use.\n"
            "- Answers summarize matching source passages; this app does not call a cloud LLM.\n"
            "- Scanned PDFs may need OCR before indexing.\n"
            "- Streamlit Cloud local files may not persist through rebuilds. Use a persistent database for important data.\n"
            "- Before a public multi-user launch, add authentication, access controls, automated tests, monitoring, and backup automation."
        )


if __name__ == "__main__":
    main()
